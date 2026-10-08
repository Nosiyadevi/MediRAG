"""Load PDF / TXT / DOCX into a list of {"page": int, "text": str}."""
import logging
from pathlib import Path

log = logging.getLogger(__name__)


def load_pages(path) -> list:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _load_pdf(path)
    if ext == ".txt":
        return _load_txt(path)
    if ext == ".docx":
        return _load_docx(path)
    raise ValueError(f"Unsupported file type: {ext}")


def _load_pdf(path: Path) -> list:
    import fitz  # PyMuPDF
    pages = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if not text.strip():
                log.warning("%s page %d has no extractable text (scanned? OCR is not supported)", path.name, i)
            pages.append({"page": i, "text": text})
    return pages


def _load_txt(path: Path) -> list:
    raw = path.read_text(encoding="utf-8", errors="replace")
    parts = raw.split("\f") if "\f" in raw else [raw]
    return [{"page": i, "text": p} for i, p in enumerate(parts, start=1)]


def _load_docx(path: Path) -> list:
    import docx
    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return [{"page": 1, "text": "\n".join(lines)}]
